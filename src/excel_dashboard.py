"""Build the Excel dashboard (excel/Retail_Orders_Dashboard.xlsx) from the pipeline outputs.

Run after src/pipeline.py:
    python src/excel_dashboard.py
"""

import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pipeline import KPI_JSON, ROOT, SQL_OUT

OUT_FILE = ROOT / "excel" / "Retail_Orders_Dashboard.xlsx"

NAVY = "2F4A6D"
STEEL = "8FA9C6"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BOLD = Font(name="Calibri", bold=True, size=11)
TITLE = Font(name="Calibri", bold=True, size=16, color=NAVY)
NOTE = Font(name="Calibri", italic=True, size=10, color="595959")
NEG_FONT = Font(color="9C0006")

MONEY = '#,##0;-#,##0'
PCT = '0.00"%"'  # SQL already returns percentages as numbers like 9.26


def nice(col):
    """'sales_share_pct' -> 'Sales share %' for readable headers and legends."""
    return col.replace("_pct", " %").replace("_", " ").capitalize()


TABLE_TITLES = set()  # left out of column auto-sizing


def result(name):
    return pd.read_csv(SQL_OUT / f"{name}.csv")


def write_table(ws, df, top, left=1, title=None, formats=None):
    """Write a DataFrame with a bold header row. Returns (first_data_row, last_data_row)."""
    formats = formats or {}
    if title:
        ws.cell(row=top, column=left, value=title).font = Font(name="Calibri", bold=True, size=12, color=NAVY)
        TABLE_TITLES.add(title)
        top += 1
    for j, col in enumerate(df.columns):
        c = ws.cell(row=top, column=left + j, value=nice(col))
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = Border(top=THIN, bottom=THIN)
        c.alignment = Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            c = ws.cell(row=top + i, column=left + j, value=value.item() if hasattr(value, "item") else value)
            fmt = formats.get(df.columns[j])
            if fmt:
                c.number_format = fmt
    last = top + len(df)
    for j in range(len(df.columns)):
        ws.cell(row=last, column=left + j).border = Border(bottom=THIN)
    return top + 1, last


def autosize(ws, min_width=10, max_width=40):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max((len(str(c.value)) for c in col
                     if c.value is not None and c.value not in TABLE_TITLES), default=0)
        ws.column_dimensions[letter].width = max(min_width, min(max_width, width + 2))


def col_range(left, offset, first, last):
    letter = get_column_letter(left + offset)
    return f"{letter}{first}:{letter}{last}"


def data_bar(ws, ref):
    ws.conditional_formatting.add(
        ref, DataBarRule(start_type="num", start_value=0, end_type="max", color=NAVY))


def negatives_red(ws, ref):
    ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=NEG_FONT))


def style_chart(chart, colors, line=False):
    # openpyxl 3.1 hides axes unless told otherwise
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.title.overlay = False
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E0E0E0"))
    for series, color in zip(chart.series, colors):
        series.invertIfNegative = False
        if line:
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = 22000
            series.smooth = False
        else:
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
    chart.legend.position = "t"
    chart.style = 2


def build():
    kpis = json.loads(KPI_JSON.read_text(encoding="utf-8"))
    wb = Workbook()

    # ---------------- data sheets ----------------
    ws_m = wb.create_sheet("Monthly")
    mom = result("q3_mom_sales_2022_vs_2023")
    mom.insert(0, "month", pd.to_datetime(mom["order_month"], format="%m").dt.strftime("%b"))
    mom = mom.drop(columns="order_month")
    m_first, m_last = write_table(ws_m, mom, 1, title="Monthly sales, 2022 vs 2023 (Q3)",
                                  formats={"sales_2022": MONEY, "sales_2023": MONEY, "growth_pct": PCT})
    negatives_red(ws_m, col_range(1, 3, m_first, m_last))
    run = result("q6_running_total_sales")
    r_first, r_last = write_table(ws_m, run, 1, left=7, title="Year-to-date running total (Q6)",
                                  formats={"sales": MONEY, "running_total_ytd": MONEY})
    autosize(ws_m)

    ws_p = wb.create_sheet("Products")
    top10 = result("q1_top10_revenue_products")
    p_first, p_last = write_table(ws_p, top10, 1, title="Top 10 products by revenue (Q1)", formats={"sales": MONEY})
    data_bar(ws_p, col_range(1, 1, p_first, p_last))
    top5 = result("q2_top5_products_per_region")
    f, l = write_table(ws_p, top5, 1, left=5, title="Top 5 products in each region (Q2)", formats={"sales": MONEY})
    data_bar(ws_p, col_range(5, 3, f, l))
    autosize(ws_p)

    ws_c = wb.create_sheet("Categories")
    share = result("q7_category_revenue_share")
    c_first, c_last = write_table(ws_c, share, 1, title="Category share of sales and profit (Q7)",
                                  formats={"sales": MONEY, "profit": MONEY,
                                           "sales_share_pct": PCT, "profit_share_pct": PCT})
    best = result("q4_best_month_per_category")
    write_table(ws_c, best, c_last + 3, title="Highest-sales month for each category (Q4)", formats={"sales": MONEY})
    sub = result("q8_subcategory_rank_within_category")
    f, l = write_table(ws_c, sub, c_last + 10, title="Sub-categories ranked by profit within category (Q8)",
                       formats={"sales": MONEY, "profit": MONEY, "margin_pct": PCT})
    data_bar(ws_c, col_range(1, 3, f, l))
    autosize(ws_c)

    ws_g = wb.create_sheet("Profit Growth")
    growth = result("q5_subcategory_profit_growth")
    g_first, g_last = write_table(ws_g, growth, 1, title="Profit growth by sub-category, 2023 vs 2022 (Q5)",
                                  formats={"profit_2022": MONEY, "profit_2023": MONEY,
                                           "profit_growth": MONEY, "growth_pct": PCT})
    negatives_red(ws_g, col_range(1, 3, g_first, g_last))
    negatives_red(ws_g, col_range(1, 4, g_first, g_last))
    autosize(ws_g)

    ws_r = wb.create_sheet("Regions & Segments")
    region = result("q9_region_profitability")
    rg_first, rg_last = write_table(ws_r, region, 1, title="Regional profitability (Q9)",
                                    formats={"sales": MONEY, "profit": MONEY, "margin_pct": PCT,
                                             "sales_growth_pct": PCT})
    negatives_red(ws_r, col_range(1, 4, rg_first, rg_last))
    data_bar(ws_r, col_range(1, 2, rg_first, rg_last))
    seg = result("q10_segment_profitability")
    f, l = write_table(ws_r, seg, rg_last + 3, title="Segment profitability (Q10)",
                       formats={"orders": "#,##0", "sales": MONEY, "profit": MONEY,
                                "margin_pct": PCT, "avg_order_value": "#,##0.00"})
    data_bar(ws_r, col_range(1, 3, f, l))
    autosize(ws_r)

    ws_d = wb.create_sheet("Discount")
    disc = result("q11_discount_impact")
    d_first, d_last = write_table(ws_d, disc, 1, title="Discount impact on profit (Q11)",
                                  formats={"order_lines": "#,##0", "sales": MONEY, "discount_given": MONEY,
                                           "profit": MONEY, "margin_pct": PCT})
    autosize(ws_d)

    # ---------------- dashboard sheet ----------------
    ws = wb.active
    ws.title = "Dashboard"
    wb.move_sheet("Dashboard", offset=-wb.index(ws))
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Retail Orders Dashboard"
    ws["A1"].font = TITLE
    ws["A2"] = (f"Order lines from {kpis['date_from']} to {kpis['date_to']}. "
                "Source: Kaggle 'Retail Orders' (ankitbansal06). Figures come from sql/analysis.sql.")
    ws["A2"].font = NOTE

    tiles = [
        ("Total sales", kpis["total_sales"], MONEY),
        ("Total profit", kpis["total_profit"], MONEY),
        ("Profit margin", kpis["profit_margin_pct"], PCT),
        ("Orders", kpis["orders"], "#,##0"),
        ("Avg order value", kpis["avg_order_value"], "#,##0.00"),
        ("Sales growth 2023 vs 2022", kpis["sales_growth_pct"], PCT),
    ]
    for i, (label, value, fmt) in enumerate(tiles):
        col = 1 + i * 2
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        top = ws.cell(row=4, column=col, value=label)
        top.font = Font(name="Calibri", size=10, color="595959")
        top.alignment = Alignment(horizontal="center")
        val = ws.cell(row=5, column=col, value=value)
        val.font = Font(name="Calibri", bold=True, size=16, color=NAVY)
        val.number_format = fmt
        val.alignment = Alignment(horizontal="center")
        for r in (4, 5):
            for c in (col, col + 1):
                ws.cell(row=r, column=c).border = Border(
                    top=THIN if r == 4 else None, bottom=THIN if r == 5 else None,
                    left=THIN if c == col else None, right=THIN if c == col + 1 else None)
    for c in range(1, 13):
        ws.column_dimensions[get_column_letter(c)].width = 11
    ws.row_dimensions[5].height = 28

    # Monthly sales line chart
    line = LineChart()
    line.title = "Monthly sales, 2022 vs 2023"
    line.height, line.width = 7.5, 16
    line.add_data(Reference(ws_m, min_col=2, max_col=3, min_row=m_first - 1, max_row=m_last), titles_from_data=True)
    line.set_categories(Reference(ws_m, min_col=1, min_row=m_first, max_row=m_last))
    line.y_axis.numFmt = "#,##0"
    style_chart(line, [STEEL, NAVY], line=True)
    ws.add_chart(line, "A7")

    # Top 10 products
    bar = BarChart()
    bar.type = "bar"
    bar.title = "Top 10 products by revenue"
    bar.height, bar.width = 7.5, 14
    bar.add_data(Reference(ws_p, min_col=2, min_row=p_first - 1, max_row=p_last), titles_from_data=True)
    bar.set_categories(Reference(ws_p, min_col=1, min_row=p_first, max_row=p_last))
    bar.x_axis.scaling.orientation = "maxMin"
    bar.y_axis.crosses = "max"
    bar.y_axis.numFmt = "#,##0"
    style_chart(bar, [NAVY])
    bar.legend = None
    ws.add_chart(bar, "I7")

    # Category share
    cat = BarChart()
    cat.title = "Category share of sales and profit (%)"
    cat.height, cat.width = 7.5, 16
    cat.add_data(Reference(ws_c, min_col=3, min_row=c_first - 1, max_row=c_last), titles_from_data=True)
    cat.add_data(Reference(ws_c, min_col=5, min_row=c_first - 1, max_row=c_last), titles_from_data=True)
    cat.set_categories(Reference(ws_c, min_col=1, min_row=c_first, max_row=c_last))
    style_chart(cat, [STEEL, NAVY])
    ws.add_chart(cat, "A23")

    # Region margin
    reg = BarChart()
    reg.title = "Profit margin by region (%)"
    reg.height, reg.width = 7.5, 14
    reg.add_data(Reference(ws_r, min_col=4, min_row=rg_first - 1, max_row=rg_last), titles_from_data=True)
    reg.set_categories(Reference(ws_r, min_col=1, min_row=rg_first, max_row=rg_last))
    style_chart(reg, [NAVY])
    reg.y_axis.scaling.min = 0
    reg.legend = None
    ws.add_chart(reg, "I23")

    # Discount vs margin
    dm = BarChart()
    dm.title = "Profit margin (%) by discount % given"
    dm.height, dm.width = 7.5, 16
    dm.add_data(Reference(ws_d, min_col=6, min_row=d_first - 1, max_row=d_last), titles_from_data=True)
    dm.set_categories(Reference(ws_d, min_col=1, min_row=d_first, max_row=d_last))
    style_chart(dm, [NAVY])
    dm.legend = None
    ws.add_chart(dm, "A39")

    # Profit growth by sub-category
    pg = BarChart()
    pg.type = "bar"
    pg.title = "Profit growth by sub-category, 2023 vs 2022"
    pg.height, pg.width = 7.5, 14
    pg.add_data(Reference(ws_g, min_col=4, min_row=g_first - 1, max_row=g_last), titles_from_data=True)
    pg.set_categories(Reference(ws_g, min_col=1, min_row=g_first, max_row=g_last))
    pg.x_axis.scaling.orientation = "maxMin"
    pg.x_axis.tickLblPos = "low"
    pg.x_axis.tickLblSkip = 1
    pg.y_axis.crosses = "max"
    pg.y_axis.numFmt = "#,##0"
    style_chart(pg, [NAVY])
    pg.legend = None
    ws.add_chart(pg, "I39")

    # ---------------- about sheet ----------------
    ws_a = wb.create_sheet("About")
    rows = [
        ("Field", "Meaning"),
        ("sale_price", "list_price minus discount. Revenue = sum of sale_price."),
        ("discount", "list_price x discount_percent / 100"),
        ("profit", "sale_price minus cost_price"),
        ("margin_pct", "profit / sales x 100"),
        ("growth_pct", "(2023 - 2022) / 2022 x 100"),
        ("", ""),
        ("Source", "Kaggle 'Retail Orders' dataset by Ankit Bansal (ankitbansal06/retail-orders)"),
        ("Built by", "Gaurav Kalamkhede. Generated by src/excel_dashboard.py from the SQL results."),
    ]
    for i, (a, b) in enumerate(rows, start=1):
        ws_a.cell(row=i, column=1, value=a).font = BOLD if i == 1 or a else Font()
        ws_a.cell(row=i, column=2, value=b)
    ws_a["A1"].fill = HEADER_FILL
    ws_a["B1"].fill = HEADER_FILL
    ws_a["A1"].font = ws_a["B1"].font = HEADER_FONT
    ws_a.column_dimensions["A"].width = 14
    ws_a.column_dimensions["B"].width = 80

    wb.properties.creator = wb.properties.lastModifiedBy = "Gaurav Kalamkhede"
    wb.properties.title = "Retail Orders Dashboard"
    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_FILE)
    print(f"Excel dashboard -> {OUT_FILE.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
