"""Build the PowerPoint deck (reports/Retail_Orders_Analysis.pptx).

Style: white slides with a navy theme, amber used sparingly as an accent, Calibri throughout.
Run after src/pipeline.py:
    python src/build_deck.py
"""

import json

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

from pipeline import CHART_DIR, KPI_JSON, ROOT, SQL_OUT

IMG_DIR = ROOT / "images"
OUT_FILE = ROOT / "reports" / "Retail_Orders_Analysis.pptx"

FONT = "Calibri"
NAVY = RGBColor(0x2F, 0x4A, 0x6D)
STEEL = RGBColor(0x8F, 0xA9, 0xC6)
PALE = RGBColor(0xEE, 0xF2, 0xF7)
AMBER = RGBColor(0xD4, 0x89, 0x2A)
TEXT = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
BORDER = RGBColor(0xD0, 0xD7, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
LEFT = Inches(0.6)
CONTENT_TOP = Inches(1.45)


# ----------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------
def rect(slide, left, top, width, height, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, value, left, top, width, height, size=16, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tf = slide.shapes.add_textbox(left, top, width, height).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = align
    p.font.name, p.font.size, p.font.bold, p.font.italic = FONT, Pt(size), bold, italic
    p.font.color.rgb = color
    return tf


def new_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    rect(slide, 0, 0, SLIDE_W, Inches(0.12), NAVY)
    text(slide, title, LEFT, Inches(0.38), SLIDE_W - 2 * LEFT, Inches(0.7), size=28, color=NAVY, bold=True)
    rect(slide, LEFT + Inches(0.03), Inches(1.08), Inches(1.3), Inches(0.06), AMBER)
    line = slide.shapes.add_connector(1, LEFT, Inches(6.95), SLIDE_W - LEFT, Inches(6.95))
    line.line.color.rgb = BORDER
    line.line.width = Pt(0.75)
    text(slide, "Retail Orders Analysis  |  Gaurav Kalamkhede", LEFT, Inches(7.0), Inches(8), Inches(0.3),
         size=10, color=GREY)
    text(slide, str(len(prs.slides)), SLIDE_W - LEFT - Inches(1), Inches(7.0), Inches(1), Inches(0.3),
         size=11, color=NAVY, bold=True, align=PP_ALIGN.RIGHT)
    return slide


def bullets(slide, items, left, top, width, height, size=16, panel=False):
    """items: list of str, or (str, level) tuples. Text in **...** at the start is bold navy."""
    if panel:
        rect(slide, left, top, width, height, PALE)
        rect(slide, left, top, Inches(0.07), height, NAVY)
        left, top = left + Inches(0.25), top + Inches(0.2)
        width, height = width - Inches(0.4), height - Inches(0.3)
    tf = slide.shapes.add_textbox(left, top, width, height).text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        value, level = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        r = p.add_run()
        r.text = ("■  " if level == 0 else "      –  ")
        r.font.name, r.font.size = FONT, Pt(size - 6 if level == 0 else size - 2)
        r.font.color.rgb = AMBER if level == 0 else GREY
        if value.startswith("**") and "**" in value[2:]:
            head, rest = value[2:].split("**", 1)
            runs = [(head, True, NAVY), (rest, False, TEXT)]
        else:
            runs = [(value, False, TEXT)]
        for t, bold, color in runs:
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = FONT, Pt(size if level == 0 else size - 2), bold
            r.font.color.rgb = color
    return tf


def picture(slide, path, left, top, max_w, max_h, border=False):
    """Add an image scaled to fit the box, keeping its shape."""
    w, h = Image.open(path).size
    scale = min(max_w / w, max_h / h)
    pic = slide.shapes.add_picture(str(path), left, top, int(w * scale), int(h * scale))
    if border:
        pic.line.color.rgb = BORDER
        pic.line.width = Pt(1)
    return pic


def caption(slide, value, left, top, width):
    text(slide, value, left, top, width, Inches(0.3), size=11, color=GREY, italic=True)


def subheading(slide, value, left, top, width):
    text(slide, value, left, top, width, Inches(0.4), size=16, color=NAVY, bold=True)


def kpi_tiles(slide, items, left, top, width, height, cols, gap=Inches(0.2)):
    """Pale tiles with a navy edge: big navy value, small grey label."""
    rows = -(-len(items) // cols)
    tile_w = (width - gap * (cols - 1)) / cols
    tile_h = (height - gap * (rows - 1)) / rows
    for i, (label, value) in enumerate(items):
        x = left + (i % cols) * (tile_w + gap)
        y = top + (i // cols) * (tile_h + gap)
        rect(slide, int(x), int(y), int(tile_w), int(tile_h), PALE)
        rect(slide, int(x), int(y), Inches(0.07), int(tile_h), NAVY)
        text(slide, value, int(x + Inches(0.25)), int(y + Inches(0.12)), int(tile_w - Inches(0.3)),
             int(tile_h * 0.55), size=26, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, label, int(x + Inches(0.25)), int(y + tile_h * 0.58), int(tile_w - Inches(0.3)),
             int(tile_h * 0.38), size=12, color=GREY)


def number_badge(slide, n, left, top, size=Inches(0.5)):
    c = rect(slide, left, top, size, size, NAVY, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = str(n)
    p.alignment = PP_ALIGN.CENTER
    p.font.name, p.font.size, p.font.bold, p.font.color.rgb = FONT, Pt(16), True, WHITE


def _no_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w="0")
        etree.SubElement(ln, qn("a:noFill"))


def table(slide, df, left, top, width, row_h=Inches(0.38), size=13, col_widths=None):
    """Navy header row with white text, light zebra rows."""
    rows, cols = len(df) + 1, len(df.columns)
    shape = slide.shapes.add_table(rows, cols, left, top, width, row_h * rows)
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    if col_widths:
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = w
    for r in range(rows):
        tbl.rows[r].height = row_h
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = str(df.columns[c] if r == 0 else df.iloc[r - 1, c])
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (PALE if r % 2 == 0 else WHITE)
            _no_border(cell)
            p = cell.text_frame.paragraphs[0]
            p.font.name, p.font.size, p.font.bold = FONT, Pt(size), r == 0
            p.font.color.rgb = WHITE if r == 0 else TEXT
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
    return shape


def fmt_money(v):
    return f"{v:,.0f}"


# ----------------------------------------------------------------------
def build():
    k = json.loads(KPI_JSON.read_text(encoding="utf-8"))
    res = {p.stem: pd.read_csv(p) for p in SQL_OUT.glob("*.csv")}
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    content_w = SLIDE_W - 2 * LEFT

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, Inches(0.45), SLIDE_H, NAVY)
    rect(slide, Inches(0.45), 0, Inches(0.07), SLIDE_H, AMBER)
    x = Inches(1.2)
    text(slide, "Retail Orders Analysis", x, Inches(1.5), Inches(11), Inches(1), size=48, color=NAVY, bold=True)
    text(slide, "End-to-end data analytics project with Python (pandas) and SQL", x, Inches(2.5),
         Inches(11), Inches(0.6), size=22, color=GREY)
    rect(slide, x + Inches(0.03), Inches(3.25), Inches(1.6), Inches(0.07), AMBER)
    text(slide, "Gaurav Kalamkhede", x, Inches(3.5), Inches(8), Inches(0.5), size=20, color=TEXT, bold=True)
    text(slide, f"{k['orders']:,} orders, {k['date_from'][:4]}-{k['date_to'][:4]}  |  github.com/gauravkal006",
         x, Inches(3.95), Inches(8), Inches(0.4), size=14, color=GREY)
    kpi_tiles(slide, [
        ("Total sales", fmt_money(k["total_sales"])),
        ("Total profit", fmt_money(k["total_profit"])),
        ("Profit margin", f"{k['profit_margin_pct']:.2f}%"),
        ("Sales growth 2023 vs 2022", f"+{k['sales_growth_pct']:.2f}%"),
    ], x, Inches(5.0), Inches(11.4), Inches(1.3), cols=4)

    # 2. Overview
    slide = new_slide(prs, "Project overview")
    text(slide, "Turn a raw retail orders file into answers about products, months, regions and profit.",
         LEFT, CONTENT_TOP, content_w, Inches(0.5), size=18, color=TEXT)
    text(slide, "Data: Kaggle 'Retail Orders' - 9,994 order lines, United States, Jan 2022 to Dec 2023.",
         LEFT, CONTENT_TOP + Inches(0.5), content_w, Inches(0.5), size=16, color=GREY)
    steps = [("Clean", "Fix missing values, rename columns, derive discount, sale price and profit in pandas"),
             ("Load", "Load the clean table into a SQLite database"),
             ("Analyse", "Answer five business questions in SQL")]
    box_w, box_h, gap = Inches(3.7), Inches(1.75), Inches(0.45)
    top = CONTENT_TOP + Inches(1.5)
    subheading(slide, "Flow from Ankit Bansal's tutorial", LEFT, top - Inches(0.45), Inches(8))
    for i, (head, body) in enumerate(steps):
        bx = LEFT + i * (box_w + gap)
        rect(slide, bx, top, box_w, box_h, PALE)
        rect(slide, bx, top, box_w, Inches(0.07), NAVY)
        number_badge(slide, i + 1, bx + Inches(0.2), top + Inches(0.25))
        text(slide, head, bx + Inches(0.85), top + Inches(0.25), box_w - Inches(1), Inches(0.5),
             size=20, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, body, bx + Inches(0.2), top + Inches(0.85), box_w - Inches(0.4), Inches(0.85), size=14)
        if i < 2:
            rect(slide, bx + box_w + Inches(0.08), top + box_h / 2 - Inches(0.15),
                         Inches(0.3), Inches(0.3), AMBER, shape=MSO_SHAPE.RIGHT_ARROW)
    subheading(slide, "Added by me", LEFT, Inches(5.0), Inches(8))
    bullets(slide, [
        "Six more SQL queries using window functions and CTEs",
        "Charts, an Excel dashboard, this deck and a PDF report",
        "A one-command pipeline (src/pipeline.py) and checks that SQL and pandas agree",
    ], LEFT, Inches(5.45), content_w, Inches(1.4), size=16)

    # 3. Cleaning
    slide = new_slide(prs, "Step 1: Cleaning the data with pandas")
    bullets(slide, [
        "**Missing values: **'Not Available' and 'unknown' read as NaN (6 ship modes).",
        "**Column names: **converted to snake_case for pandas and SQL.",
        "**New columns: **discount, sale_price and profit.",
        "**Dates: **order_date converted from text to a date.",
        "**Dropped: **list_price and cost_price (already captured).",
        "**Result: **9,994 rows, no duplicate orders, no rows lost.",
    ], LEFT, CONTENT_TOP, Inches(5.1), Inches(4.4), size=16, panel=True)
    pic = picture(slide, IMG_DIR / "nb_02_derive_columns.png", Inches(6.0), CONTENT_TOP, Inches(6.7), Inches(5.0), border=True)
    caption(slide, "Notebook: deriving discount, sale_price and profit", Inches(6.0), pic.top + pic.height + Inches(0.1), Inches(6.7))

    # 4. Cleaning checks
    slide = new_slide(prs, "Checks after cleaning")
    pic = picture(slide, IMG_DIR / "nb_03_cleaning_checks.png", LEFT, CONTENT_TOP, Inches(7.4), Inches(5.0), border=True)
    caption(slide, "Notebook: checks before and after cleaning", LEFT, pic.top + pic.height + Inches(0.1), Inches(7))
    bullets(slide, [
        "No rows lost, every order_id unique.",
        "507 lines have a price and cost of 0. They add nothing, so they are kept.",
        "**3,997 lines (40%) lose money.**",
        ("3,955 have list price = cost, so any discount makes a loss", 1),
        ("Total loss is 3,151 (1.5% of profit)", 1),
    ], Inches(8.3), CONTENT_TOP, Inches(4.4), Inches(3.5), size=16, panel=True)

    # 5. KPIs
    slide = new_slide(prs, "Headline numbers")
    kpi_tiles(slide, [
        ("Total sales", fmt_money(k["total_sales"])),
        ("Total profit", fmt_money(k["total_profit"])),
        ("Profit margin", f"{k['profit_margin_pct']:.2f}%"),
        ("Orders", f"{k['orders']:,}"),
        ("Average order value", f"{k['avg_order_value']:,.2f}"),
        ("Total discount given", fmt_money(k["total_discount"])),
        ("Sales growth 2023 vs 2022", f"+{k['sales_growth_pct']:.2f}%"),
        ("Profit growth 2023 vs 2022", f"+{k['profit_growth_pct']:.2f}%"),
    ], LEFT, CONTENT_TOP, content_w, Inches(3.0), cols=4)
    bullets(slide, [
        f"Sales rose from {fmt_money(k['sales_2022'])} in 2022 to {fmt_money(k['sales_2023'])} in 2023.",
        "Profit grew slightly slower than sales, so the margin dipped from 9.28% to 9.24%.",
        "About 1 in every 11 dollars of sales is profit.",
    ], LEFT, Inches(4.75), content_w, Inches(1.9), size=17, panel=True)

    # 6. Q1
    slide = new_slide(prs, "Q1. Top 10 products by revenue")
    picture(slide, CHART_DIR / "02_top10_products.png", LEFT, CONTENT_TOP, Inches(8.0), Inches(5.3))
    bullets(slide, [
        f"**{k['top_product']}** (a copier) earns {fmt_money(k['top_product_sales'])}, more than twice the second product.",
        "It is 2.7% of all sales on its own.",
        "Four of the top ten are binders; two are machines.",
    ], Inches(8.8), CONTENT_TOP + Inches(0.2), Inches(3.95), Inches(2.9), size=16, panel=True)

    # 7. Q2 and Q4
    slide = new_slide(prs, "Q2 and Q4. Leaders by region and best month by category")
    subheading(slide, "Q2: No. 1 product in each region", LEFT, CONTENT_TOP, Inches(6))
    top1 = res["q2_top5_products_per_region"].query("rank_in_region == 1")[["region", "product_id", "sales"]]
    top1 = top1.assign(sales=top1["sales"].map(fmt_money))
    top1.columns = ["Region", "No. 1 product", "Sales"]
    table(slide, top1, LEFT, CONTENT_TOP + Inches(0.5), Inches(5.6), size=15,
          col_widths=[Inches(1.6), Inches(2.5), Inches(1.5)])
    subheading(slide, "Q4: Highest-sales month by category", Inches(7.0), CONTENT_TOP, Inches(6))
    best = res["q4_best_month_per_category"].copy()
    best["sales"] = best["sales"].map(fmt_money)
    best.columns = ["Category", "Best month", "Sales"]
    table(slide, best, Inches(7.0), CONTENT_TOP + Inches(0.5), Inches(5.6), size=15,
          col_widths=[Inches(2.4), Inches(1.7), Inches(1.5)])
    bullets(slide, [
        "The same copier leads three of the four regions. The South is led by a machine.",
        "October appears twice: it is a strong month for Furniture and Technology.",
        "Full top-5 lists for every region are in sql/analysis.sql (Q2).",
    ], LEFT, Inches(4.6), content_w, Inches(1.55), size=16, panel=True)

    # 8. Q3
    slide = new_slide(prs, "Q3. Month-over-month sales, 2022 vs 2023")
    picture(slide, CHART_DIR / "01_monthly_sales_2022_vs_2023.png", LEFT, CONTENT_TOP, Inches(8.3), Inches(5.3))
    bullets(slide, [
        "2023 beat 2022 in 7 of 12 months.",
        "**Biggest gains: **February (+42%) and April (+17%).",
        "**Biggest falls: **June (-27%) and August (-16%).",
        "October was a top-two month in both years.",
    ], Inches(9.1), CONTENT_TOP + Inches(0.2), Inches(3.65), Inches(3.3), size=16, panel=True)

    # 9. Q5
    g = res["q5_subcategory_profit_growth"].iloc[0]
    slide = new_slide(prs, "Q5. Which sub-category grew profit the most?")
    picture(slide, CHART_DIR / "03_subcategory_profit_growth.png", LEFT, CONTENT_TOP, Inches(7.6), Inches(5.4))
    bullets(slide, [
        f"**{g['sub_category']}: **profit up from {fmt_money(g['profit_2022'])} to {fmt_money(g['profit_2023'])} (+{g['growth_pct']:.0f}%).",
        "Phones, Binders and Storage also grew strongly.",
        "Copiers, Appliances and Tables lost the most, about 7.6k together.",
        "9 sub-categories grew, 8 declined.",
    ], Inches(8.5), CONTENT_TOP + Inches(0.2), Inches(4.25), Inches(3.3), size=16, panel=True)

    # 10. SQL techniques
    slide = new_slide(prs, "Step 2: SQL in SQLite - techniques used")
    bullets(slide, [
        "**CTEs (WITH) **to break queries into steps",
        "**ROW_NUMBER, RANK, DENSE_RANK **for top-N per group",
        "**SUM() OVER (PARTITION BY ... ORDER BY ...) **for running totals",
        "**SUM(SUM(x)) OVER () **for share of total",
        "**CASE WHEN **to pivot 2022 and 2023 into columns",
        "**strftime() **for year and month in SQLite",
    ], LEFT, CONTENT_TOP, Inches(5.5), Inches(3.9), size=16, panel=True)
    pic = picture(slide, IMG_DIR / "nb_07_sql_q9_region.png", Inches(6.4), CONTENT_TOP, Inches(6.3), Inches(4.8), border=True)
    caption(slide, "Notebook: Q9 regional profitability query result", Inches(6.4), pic.top + pic.height + Inches(0.1), Inches(6.3))

    # 11. Category and region
    slide = new_slide(prs, "Categories and regions")
    half = (content_w - Inches(0.4)) / 2
    picture(slide, CHART_DIR / "04_category_share.png", LEFT, CONTENT_TOP, half, Inches(3.7))
    picture(slide, CHART_DIR / "05_region_margin.png", LEFT + half + Inches(0.4), CONTENT_TOP, half, Inches(3.7))
    reg = res["q9_region_profitability"].set_index("region")
    bullets(slide, [
        "Technology is the largest category (36% of sales, 37% of profit). Margins are similar across categories.",
        f"South has the best margin ({reg.loc['South', 'margin_pct']:.2f}%) but the lowest sales.",
        f"West sells the most, but its sales fell {abs(reg.loc['West', 'sales_growth_pct']):.1f}% in 2023 while the other regions grew.",
    ], LEFT, Inches(5.05), content_w, Inches(1.75), size=15, panel=True)

    # 12. Discount
    slide = new_slide(prs, "Discounts reduce margin")
    picture(slide, CHART_DIR / "06_discount_vs_margin.png", LEFT, CONTENT_TOP, Inches(7.6), Inches(5.2))
    d = res["q11_discount_impact"].set_index("discount_percent")
    bullets(slide, [
        f"Margin falls from {d.loc[2, 'margin_pct']:.2f}% at a 2% discount to {d.loc[5, 'margin_pct']:.2f}% at 5%.",
        "Each discount level has about 2,500 lines, so this is not a few large orders.",
        f"The 4% and 5% levels gave {fmt_money(d.loc[4, 'discount_given'] + d.loc[5, 'discount_given'])} of discount.",
        "Items with no markup still get discounts and make a loss.",
    ], Inches(8.5), CONTENT_TOP + Inches(0.2), Inches(4.25), Inches(3.6), size=16, panel=True)

    # 13. Excel dashboard
    slide = new_slide(prs, "Excel dashboard")
    pic = picture(slide, IMG_DIR / "excel_dashboard.png", LEFT, CONTENT_TOP, content_w, Inches(5.0), border=True)
    pic.left = int((SLIDE_W - pic.width) / 2)
    caption(slide, "excel/Retail_Orders_Dashboard.xlsx - KPI cells, native Excel charts, and data sheets with conditional formatting",
            pic.left, pic.top + pic.height + Inches(0.08), content_w)

    # 14. Recommendations
    slide = new_slide(prs, "Recommendations")
    recs = [
        ("Stop discounting items sold at cost", "This removes almost all loss-making lines."),
        ("Review the 4-5% discount levels", "They cost about 51k and give the lowest margins."),
        ("Look into the West", "The biggest region, but the only one that shrank in 2023."),
        ("Protect the top copier", "One product brings 2.7% of all sales."),
        ("Check Copiers, Appliances and Tables", "Find out why their profit fell: price, cost or mix."),
        ("Plan for busy months", "October is strong every year; February and April were strong in 2023."),
    ]
    col_w, row_h = (content_w - Inches(0.4)) / 2, Inches(1.6)
    for i, (head, body) in enumerate(recs):
        bx = LEFT + (i % 2) * (col_w + Inches(0.4))
        by = CONTENT_TOP + (i // 2) * (row_h + Inches(0.15))
        rect(slide, int(bx), int(by), int(col_w), int(row_h), PALE)
        number_badge(slide, i + 1, int(bx + Inches(0.25)), int(by + Inches(0.3)), size=Inches(0.6))
        text(slide, head, int(bx + Inches(1.1)), int(by + Inches(0.25)), int(col_w - Inches(1.3)), Inches(0.5),
             size=18, color=NAVY, bold=True)
        text(slide, body, int(bx + Inches(1.1)), int(by + Inches(0.75)), int(col_w - Inches(1.3)), Inches(0.7),
             size=15, color=TEXT)

    # 15. Tools and acknowledgement
    slide = new_slide(prs, "Tools and acknowledgement")
    bullets(slide, [
        "**Python: **pandas (cleaning), matplotlib (charts), openpyxl (Excel), python-pptx (slides), Playwright (PDF)",
        "**SQL: **SQLite - 11 queries in sql/analysis.sql",
        "**Notebook: **notebooks/retail_orders_analysis.ipynb, executed with outputs",
        "**Reproducible: **python src/pipeline.py rebuilds the clean data, SQL results, KPIs and charts",
    ], LEFT, CONTENT_TOP, content_w, Inches(2.6), size=17)
    subheading(slide, "Acknowledgement", LEFT, Inches(3.8), Inches(6))
    bullets(slide, [
        "The flow and the five core questions follow Ankit Bansal's YouTube tutorial 'End to End Data Analytics "
        "Project (Python + SQL)' and his Kaggle dataset. The code was written fresh, and the extended analysis, "
        "dashboard, report and slides were added.",
    ], LEFT, Inches(4.25), content_w, Inches(1.15), size=16, panel=True)

    props = prs.core_properties
    props.author = props.last_modified_by = "Gaurav Kalamkhede"
    props.title = "Retail Orders Analysis"
    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_FILE)
    print(f"Deck -> {OUT_FILE.relative_to(ROOT)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
